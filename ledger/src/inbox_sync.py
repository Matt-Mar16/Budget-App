"""
inbox_sync.py — phone-entry bridge for THE LEDGER.

Data flow: build_inbox_workbook() generates Inbox.xlsx (a Transactions
table plus Accounts/Categories reference sheets exported live from the
profile's database, used as dropdown sources) somewhere phone-synced —
entries typed there on a phone reach the laptop via that sync, not via
this code. sync_inbox() is the other half: run on the laptop, it reads
whatever rows were typed, turns them into real transactions/transfers,
and clears the sheet so the next sync doesn't re-import old rows.

Safety order in sync_inbox() matters and is deliberate: every row is
validated against known accounts/categories BEFORE anything is written
anywhere (validate-before-mutate), then everything read is backed up to
a CSV BEFORE the sheet is cleared (backup-before-clear) — so a crash or
a bad row can't silently lose data, and a botched sync is always
recoverable from the backup file. There's no automatic dedup: entering
the same row twice creates two transactions, on purpose (a content-based
fingerprint would also silently drop a legitimate same-day duplicate,
e.g. two identical coffees).
"""

import csv
import datetime
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

TRANSACTIONS_SHEET = "Transactions"
TRANSACTIONS_TABLE = "InboxTable"
COLUMNS = ["Date", "Account", "Category", "Payee", "Amount", "Currency", "Note"]
VALIDATION_ROWS = 100
CURRENCY_EXCHANGE_CATEGORY = "Currency Exchange"


def build_inbox_workbook(db, path):
    accounts = db.list_accounts()
    categories = db.list_categories()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TRANSACTIONS_SHEET
    ws.append(COLUMNS)

    last_col = chr(ord("A") + len(COLUMNS) - 1)
    table = Table(displayName=TRANSACTIONS_TABLE, ref=f"A1:{last_col}1")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    accounts_ws = wb.create_sheet("Accounts")
    accounts_ws.append(["Name"])
    for acc in accounts:
        accounts_ws.append([acc["name"]])

    categories_ws = wb.create_sheet("Categories")
    categories_ws.append(["Name"])
    for cat in categories:
        categories_ws.append([cat["name"]])
    categories_ws.append([CURRENCY_EXCHANGE_CATEGORY])

    last_row = 1 + VALIDATION_ROWS

    account_dv = DataValidation(
        type="list", formula1=f"=Accounts!$A$2:$A${max(2, len(accounts) + 1)}", allow_blank=True
    )
    ws.add_data_validation(account_dv)
    account_dv.add(f"B2:B{last_row}")

    category_dv = DataValidation(
        type="list", formula1=f"=Categories!$A$2:$A${max(2, len(categories) + 2)}", allow_blank=True
    )
    ws.add_data_validation(category_dv)
    category_dv.add(f"C2:C{last_row}")

    known_currencies = sorted(db.get_fx_rates().keys()) or ["GBP"]
    currency_dv = DataValidation(type="list", formula1=f'"{",".join(known_currencies)}"', allow_blank=True)
    ws.add_data_validation(currency_dv)
    currency_dv.add(f"F2:F{last_row}")

    wb.save(path)


def _read_data_rows(ws):
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if all(value is None for value in row):
            continue
        rows.append(dict(zip(COLUMNS, row)))
    return rows


def _clear_data_rows(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    if TRANSACTIONS_TABLE in ws.tables:
        last_col = chr(ord("A") + len(COLUMNS) - 1)
        ws.tables[TRANSACTIONS_TABLE].ref = f"A1:{last_col}1"


def _format_date(value):
    return value.strftime("%Y-%m-%d") if hasattr(value, "strftime") else str(value)


def _validate_exchange_pairs(exchange_rows):
    # A currency exchange is entered as two separate inbox rows (one leg
    # spending currency A, one leg receiving currency B) sharing the same
    # Date -- group by date to find each pair before turning them into one
    # linked transfer_between_accounts() call below.
    by_date = {}
    for row in exchange_rows:
        by_date.setdefault(_format_date(row["Date"]), []).append(row)

    for date, rows_for_date in by_date.items():
        if len(rows_for_date) != 2:
            raise ValueError(
                f"Currency Exchange rows on {date} aren't paired — expected exactly 2 "
                f"(one outgoing, one incoming), found {len(rows_for_date)}"
            )
        amounts = [r["Amount"] for r in rows_for_date]
        if not (amounts[0] < 0 < amounts[1] or amounts[1] < 0 < amounts[0]):
            raise ValueError(f"Currency Exchange rows on {date} must have one negative and one positive Amount")

    return by_date


def sync_inbox(excel_path, db, backups_dir, now=None):
    now = now or datetime.datetime.now()

    wb = openpyxl.load_workbook(excel_path)
    ws = wb[TRANSACTIONS_SHEET]

    rows = _read_data_rows(ws)
    if not rows:
        return 0

    accounts_by_name = {acc["name"]: acc["id"] for acc in db.list_accounts()}
    categories_by_name = {cat["name"]: cat["id"] for cat in db.list_categories()}

    exchange_rows = [r for r in rows if r["Category"] == CURRENCY_EXCHANGE_CATEGORY]
    normal_rows = [r for r in rows if r["Category"] != CURRENCY_EXCHANGE_CATEGORY]

    # Validate every row before writing anything -- one bad row (typo'd
    # account/category name, an unpaired exchange) aborts the whole sync
    # instead of partially importing and leaving the sheet in a confusing
    # half-cleared state.
    for row in rows:
        if row["Account"] not in accounts_by_name:
            raise ValueError(f"Unknown account {row['Account']!r} — not found in THE LEDGER's Accounts")
    for row in normal_rows:
        if row["Category"] not in categories_by_name:
            raise ValueError(f"Unknown category {row['Category']!r} — not found in THE LEDGER's Categories")

    exchange_pairs_by_date = _validate_exchange_pairs(exchange_rows)

    # Back up exactly what was read before touching the database or the
    # sheet, so an interrupted sync (crash, killed process) can always be
    # recovered by hand from this CSV.
    backups_dir = Path(backups_dir)
    backups_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backups_dir / f"pushed_{now.strftime('%Y-%m-%d_%H%M%S')}.csv"
    with open(backup_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    for row in normal_rows:
        db.add_transaction(
            date=_format_date(row["Date"]),
            payee=row["Payee"] or "",
            category_id=categories_by_name[row["Category"]],
            amount=row["Amount"],
            currency=row["Currency"],
            note=row["Note"] or "",
            account_id=accounts_by_name[row["Account"]],
        )

    for date, (row_a, row_b) in exchange_pairs_by_date.items():
        from_row, to_row = (row_a, row_b) if row_a["Amount"] < 0 else (row_b, row_a)
        db.transfer_between_accounts(
            accounts_by_name[from_row["Account"]],
            accounts_by_name[to_row["Account"]],
            abs(from_row["Amount"]),
            to_amount=abs(to_row["Amount"]),
            date=date,
            note=from_row["Payee"] or to_row["Payee"] or "",
        )

    _clear_data_rows(ws)
    wb.save(excel_path)

    return len(rows)
