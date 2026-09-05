import datetime

import openpyxl
import pytest

from finance_core import Database
from inbox_sync import build_inbox_workbook, sync_inbox


def _db(tmp_path):
    return Database(str(tmp_path / "test.db"))


def test_build_inbox_workbook_lists_current_accounts_and_categories(tmp_path):
    db = _db(tmp_path)
    db.add_account("UK Bank", "asset", 1200.0, currency="GBP")
    db.add_category("Groceries", "need")

    workbook_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, workbook_path)
    db.close()

    wb = openpyxl.load_workbook(workbook_path)
    account_names = [row[0] for row in wb["Accounts"].iter_rows(min_row=2, values_only=True)]
    category_names = [row[0] for row in wb["Categories"].iter_rows(min_row=2, values_only=True)]

    assert "UK Bank" in account_names
    assert "Groceries" in category_names
    assert wb["Transactions"].max_row == 1


def test_sync_inbox_creates_a_transaction_and_clears_the_sheet(tmp_path):
    db = _db(tmp_path)
    db.add_account("UK Bank", "asset", 1200.0, currency="GBP")
    db.add_category("Groceries", "need")

    excel_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.append(["2026-09-03", "UK Bank", "Groceries", "Tesco", -42.50, "GBP", "weekly shop"])
    wb.save(excel_path)

    synced_count = sync_inbox(excel_path, db, tmp_path / "backups")

    assert synced_count == 1

    txns = db.transactions_in_month(2026, 9)
    assert len(txns) == 1
    assert txns[0]["payee"] == "Tesco"
    assert txns[0]["amount"] == pytest.approx(-42.50)
    assert txns[0]["account_name"] == "UK Bank"
    assert txns[0]["category_name"] == "Groceries"

    ws_after = openpyxl.load_workbook(excel_path)["Transactions"]
    assert ws_after.max_row == 1

    db.close()


def test_sync_inbox_updates_account_balance(tmp_path):
    db = _db(tmp_path)
    db.add_account("UK Bank", "asset", 1200.0, currency="GBP")
    db.add_category("Groceries", "need")
    account_id = db.list_accounts()[0]["id"]

    excel_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.append(["2026-09-03", "UK Bank", "Groceries", "Tesco", -42.50, "GBP", ""])
    wb.save(excel_path)

    sync_inbox(excel_path, db, tmp_path / "backups")

    updated = db.get_account(account_id)
    assert updated["balance"] == pytest.approx(1200.0 - 42.50)

    db.close()


def test_sync_inbox_raises_a_clear_error_for_an_unknown_account_name(tmp_path):
    db = _db(tmp_path)

    excel_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.append(["2026-09-03", "Typo Bank", "Groceries", "Tesco", -42.50, "GBP", ""])
    wb.save(excel_path)

    with pytest.raises(ValueError, match="Typo Bank"):
        sync_inbox(excel_path, db, tmp_path / "backups")

    db.close()


def test_sync_inbox_writes_a_backup_of_exactly_what_it_pushed(tmp_path):
    db = _db(tmp_path)
    db.add_account("UK Bank", "asset", 1200.0, currency="GBP")
    db.add_category("Groceries", "need")

    excel_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.append(["2026-09-03", "UK Bank", "Groceries", "Tesco", -42.50, "GBP", ""])
    wb.save(excel_path)

    backups_dir = tmp_path / "backups"
    sync_inbox(excel_path, db, backups_dir, now=datetime.datetime(2026, 9, 12, 14, 30, 0))

    backup_files = list(backups_dir.glob("*.csv"))
    assert len(backup_files) == 1
    assert "2026-09-12" in backup_files[0].name

    db.close()


def test_sync_inbox_logs_a_currency_exchange_as_a_linked_transfer(tmp_path):
    db = _db(tmp_path)
    db.add_account("UK Bank", "asset", 1200.0, currency="GBP")
    db.add_account("CAD Chequing", "asset", 0.0, currency="CAD")
    accounts = {a["name"]: a["id"] for a in db.list_accounts()}

    excel_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.append(["2026-09-05", "UK Bank", "Currency Exchange", "GBP exchanged", -500.0, "GBP", ""])
    ws.append(["2026-09-05", "CAD Chequing", "Currency Exchange", "CAD received", 860.0, "CAD", ""])
    wb.save(excel_path)

    synced_count = sync_inbox(excel_path, db, tmp_path / "backups")

    assert synced_count == 2
    assert db.get_account(accounts["UK Bank"])["balance"] == pytest.approx(1200.0 - 500.0)
    assert db.get_account(accounts["CAD Chequing"])["balance"] == pytest.approx(860.0)

    transfers = db.list_transfers()
    assert len(transfers) == 1
    assert transfers[0]["historical_rate"] == pytest.approx(860.0 / 500.0)

    ws_after = openpyxl.load_workbook(excel_path)["Transactions"]
    assert ws_after.max_row == 1

    db.close()


def test_sync_inbox_raises_a_clear_error_for_an_unpaired_currency_exchange_row(tmp_path):
    db = _db(tmp_path)
    db.add_account("UK Bank", "asset", 1200.0, currency="GBP")

    excel_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.append(["2026-09-05", "UK Bank", "Currency Exchange", "GBP exchanged", -500.0, "GBP", ""])
    wb.save(excel_path)

    with pytest.raises(ValueError, match="2026-09-05"):
        sync_inbox(excel_path, db, tmp_path / "backups")

    db.close()


def test_sync_inbox_with_no_new_rows_returns_zero(tmp_path):
    db = _db(tmp_path)
    excel_path = tmp_path / "Inbox.xlsx"
    build_inbox_workbook(db, excel_path)

    synced_count = sync_inbox(excel_path, db, tmp_path / "backups")

    assert synced_count == 0
    db.close()
