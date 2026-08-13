import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

ACCOUNTS = [
    ("UK Bank", "GBP", 1200.00, "2026-09-01"),
    ("CAD Chequing", "CAD", 0.00, "2026-09-01"),
    ("CAD Credit Card", "CAD", 0.00, "2026-09-01"),
]

CATEGORIES = [
    ("Groceries", "CAD", 400),
    ("Rent", "CAD", 900),
    ("Subscriptions", "GBP", 30),
    ("Income", "GBP", None),
]

CURRENCY_EXCHANGE_CATEGORY = "Currency Exchange"

VALIDATION_ROWS = 100


def _build_accounts_sheet(wb):
    ws = wb.active
    ws.title = "Accounts"
    ws.append(["Account", "Currency", "Starting balance", "Starting date"])
    for row in ACCOUNTS:
        ws.append(list(row))
    return ws


def _build_categories_sheet(wb):
    ws = wb.create_sheet("Categories")
    ws.append(["Category", "Currency", "Monthly budget", "Label"])
    row_num = 2
    for category, currency, budget in CATEGORIES:
        # "Label" is a formula (e.g. =A2&" ("&B2&")") so it stays in sync if
        # the category/currency cells are ever hand-edited in Excel; this is
        # the column the Transactions sheet's Category dropdown reads from.
        ws.append([category, currency, budget, f'=A{row_num}&" ("&B{row_num}&")"'])
        row_num += 1
    # Currency Exchange has no currency of its own, so it gets a literal
    # string here instead of the formula above — the formula would otherwise
    # render as "Currency Exchange ()" with the Currency cell left blank.
    ws.append([CURRENCY_EXCHANGE_CATEGORY, None, None, CURRENCY_EXCHANGE_CATEGORY])
    return ws


def _build_transactions_sheet(wb):
    ws = wb.create_sheet("Transactions")
    ws.append(["Date", "Account", "Category", "Description", "Amount", "Currency", "Exchange Rate"])
    # This example row must stay: Excel's table "calculated column" auto-fill
    # only propagates the Currency formula to new rows the user adds once at
    # least one existing row already carries it. Column order in this literal
    # list is Date, Account, Category, Description, Amount, Currency, Exchange
    # Rate — the Currency cell (index 5) is the XLOOKUP formula itself.
    ws.append([
        "2026-09-01",
        ACCOUNTS[0][0],
        f"Groceries ({CATEGORIES[0][1]})",
        "Example — delete this row before logging real transactions",
        0,
        "=XLOOKUP(C2,Categories!$D:$D,Categories!$B:$B,\"\")",
        None,
    ])

    # Dropdown validation ranges are pre-extended a fixed number of rows below
    # the header so a user typing new rows keeps getting the Account/Category
    # dropdowns without needing to re-run this script.
    last_row = 1 + VALIDATION_ROWS
    table = Table(displayName="TransactionsTable", ref=f"A1:G{max(2, ws.max_row)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    account_dv = DataValidation(type="list", formula1=f"=Accounts!$A$2:$A${last_row}", allow_blank=True)
    ws.add_data_validation(account_dv)
    account_dv.add(f"B2:B{last_row}")

    category_dv = DataValidation(type="list", formula1=f"=Categories!$D$2:$D${last_row}", allow_blank=True)
    ws.add_data_validation(category_dv)
    category_dv.add(f"C2:C{last_row}")

    currency_dv = DataValidation(type="list", formula1='"GBP,CAD"', allow_blank=True)
    ws.add_data_validation(currency_dv)
    currency_dv.add(f"F2:F{last_row}")

    return ws


def build_template(path):
    wb = openpyxl.Workbook()
    _build_accounts_sheet(wb)
    _build_categories_sheet(wb)
    _build_transactions_sheet(wb)
    wb.save(path)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Generate a fresh Trip Budget Excel workbook.")
    parser.add_argument("--out", default="Data.xlsx", help="Output path for the workbook")
    args = parser.parse_args()

    build_template(args.out)
    print(f"Wrote {args.out}")
