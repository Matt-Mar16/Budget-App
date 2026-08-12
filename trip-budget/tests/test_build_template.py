import openpyxl

from build_template import build_template


def test_build_template_creates_three_sheets_with_expected_headers(tmp_path):
    path = tmp_path / "Data.xlsx"

    build_template(path)

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Accounts", "Categories", "Transactions"]

    accounts = wb["Accounts"]
    assert [cell.value for cell in accounts[1]] == ["Account", "Currency", "Starting balance", "Starting date"]

    categories = wb["Categories"]
    assert [cell.value for cell in categories[1]] == ["Category", "Currency", "Monthly budget", "Label"]

    transactions = wb["Transactions"]
    assert [cell.value for cell in transactions[1]] == [
        "Date", "Account", "Category", "Description", "Amount", "Currency", "Exchange Rate"
    ]


def test_build_template_categories_label_column_is_formula_except_currency_exchange(tmp_path):
    path = tmp_path / "Data.xlsx"

    build_template(path)

    wb = openpyxl.load_workbook(path)
    categories = wb["Categories"]
    rows = {row[0].value: row for row in categories.iter_rows(min_row=2)}

    groceries_row = rows["Groceries"]
    assert groceries_row[3].value == f'=A{groceries_row[0].row}&" ("&B{groceries_row[0].row}&")"'

    exchange_row = rows["Currency Exchange"]
    assert exchange_row[3].value == "Currency Exchange"


def test_build_template_transactions_is_a_table_with_dropdowns_and_currency_formula(tmp_path):
    path = tmp_path / "Data.xlsx"

    build_template(path)

    wb = openpyxl.load_workbook(path)
    transactions = wb["Transactions"]

    assert "TransactionsTable" in transactions.tables

    validations = transactions.data_validations.dataValidation
    account_dv = next(dv for dv in validations if "Accounts" in dv.formula1)
    category_dv = next(dv for dv in validations if "Categories" in dv.formula1)
    currency_dv = next(dv for dv in validations if dv.formula1 == '"GBP,CAD"')
    assert account_dv.type == "list"
    assert category_dv.type == "list"
    assert currency_dv.type == "list"

    example_row = transactions[2]
    assert example_row[5].value.startswith("=XLOOKUP(")
