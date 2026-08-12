import openpyxl
import pandas as pd

from build_template import build_template
from dashboard import run_dashboard


def test_run_dashboard_syncs_new_rows_then_reports_and_clears_the_sheet(tmp_path, capsys):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    charts_dir = tmp_path / "charts"
    backups_dir = tmp_path / "backups"

    build_template(excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.delete_rows(2, 1)
    ws.append(["2026-09-03", "CAD Chequing", "Groceries (CAD)", "Superstore", -62.40, "CAD", None])
    wb.save(excel_path)

    run_dashboard(excel_path, log_path, charts_dir, backups_dir)

    log_df = pd.read_csv(log_path)
    assert len(log_df) == 1
    assert log_df.loc[0, "Description"] == "Superstore"

    ws_after = openpyxl.load_workbook(excel_path)["Transactions"]
    assert ws_after.max_row == 1

    assert (charts_dir / "money_flow.html").exists()

    captured = capsys.readouterr()
    assert "budget vs actual" in captured.out.lower()


def test_run_dashboard_reports_with_no_new_rows_since_last_sync(tmp_path, capsys):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    charts_dir = tmp_path / "charts"
    backups_dir = tmp_path / "backups"

    build_template(excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.delete_rows(2, 1)
    wb.save(excel_path)

    run_dashboard(excel_path, log_path, charts_dir, backups_dir)

    captured = capsys.readouterr()
    assert "account balances" in captured.out.lower()


def test_run_dashboard_reports_on_the_given_month(tmp_path, capsys):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    charts_dir = tmp_path / "charts"
    backups_dir = tmp_path / "backups"

    build_template(excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.delete_rows(2, 1)
    ws.append(["2026-09-03", "CAD Chequing", "Groceries (CAD)", "Superstore", -62.40, "CAD", None])
    wb.save(excel_path)

    report = run_dashboard(excel_path, log_path, charts_dir, backups_dir, month=(2026, 9))

    assert report["month"] == (2026, 9)
    captured = capsys.readouterr()
    assert "2026-09" in captured.out
