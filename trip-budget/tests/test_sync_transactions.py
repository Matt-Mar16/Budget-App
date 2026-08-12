import datetime

import openpyxl
import pandas as pd

from build_template import build_template
from sync_transactions import sync


def _bare_transactions_workbook(path, data_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Date", "Account", "Category", "Description", "Amount", "Currency", "Exchange Rate"])
    for row in data_rows:
        ws.append(row)
    wb.save(path)


def test_sync_moves_rows_into_log_and_clears_the_sheet_back_to_headers(tmp_path):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    backups_dir = tmp_path / "backups"
    _bare_transactions_workbook(excel_path, [
        ["2026-09-03", "CAD Chequing", "Groceries (CAD)", "Superstore", -62.40, "CAD", None],
    ])

    synced_count = sync(excel_path, log_path, backups_dir)

    assert synced_count == 1

    log_df = pd.read_csv(log_path)
    assert len(log_df) == 1
    assert log_df.loc[0, "Description"] == "Superstore"

    ws_after = openpyxl.load_workbook(excel_path)["Transactions"]
    assert ws_after.max_row == 1


def test_sync_with_no_new_rows_returns_zero_and_never_creates_the_log(tmp_path):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    backups_dir = tmp_path / "backups"
    _bare_transactions_workbook(excel_path, [])

    synced_count = sync(excel_path, log_path, backups_dir)

    assert synced_count == 0
    assert not log_path.exists()


def test_sync_writes_a_timestamped_backup_of_exactly_what_it_pushed(tmp_path):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    backups_dir = tmp_path / "backups"
    _bare_transactions_workbook(excel_path, [
        ["2026-09-03", "CAD Chequing", "Groceries (CAD)", "Superstore", -62.40, "CAD", None],
    ])

    sync(excel_path, log_path, backups_dir, now=datetime.datetime(2026, 9, 12, 14, 30, 0))

    backup_files = list(backups_dir.glob("*.csv"))
    assert len(backup_files) == 1
    assert "2026-09-12" in backup_files[0].name
    backup_df = pd.read_csv(backup_files[0])
    assert len(backup_df) == 1
    assert backup_df.loc[0, "Description"] == "Superstore"


def test_sync_appends_to_an_existing_log_rather_than_overwriting_it(tmp_path):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    backups_dir = tmp_path / "backups"

    _bare_transactions_workbook(excel_path, [
        ["2026-09-03", "CAD Chequing", "Groceries (CAD)", "Superstore", -62.40, "CAD", None],
    ])
    sync(excel_path, log_path, backups_dir, now=datetime.datetime(2026, 9, 3, 9, 0, 0))

    _bare_transactions_workbook(excel_path, [
        ["2026-09-10", "CAD Chequing", "Rent (CAD)", "Rent", -900.0, "CAD", None],
    ])
    sync(excel_path, log_path, backups_dir, now=datetime.datetime(2026, 9, 10, 9, 0, 0))

    log_df = pd.read_csv(log_path)
    assert len(log_df) == 2
    assert list(log_df["Description"]) == ["Superstore", "Rent"]


def test_sync_shrinks_the_real_excel_table_ref_back_to_the_header_row(tmp_path):
    excel_path = tmp_path / "Data.xlsx"
    log_path = tmp_path / "transactions_log.csv"
    backups_dir = tmp_path / "backups"

    build_template(excel_path)
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Transactions"]
    ws.append(["2026-09-05", "CAD Chequing", "Groceries (CAD)", "Superstore", -62.40, "CAD", None])
    wb.save(excel_path)

    sync(excel_path, log_path, backups_dir)

    wb_after = openpyxl.load_workbook(excel_path)
    ws_after = wb_after["Transactions"]
    assert ws_after.max_row == 1
    assert ws_after.tables["TransactionsTable"].ref == "A1:G1"
